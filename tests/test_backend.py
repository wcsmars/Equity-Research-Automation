"""Offline regression tests for the FastAPI backend.

Every test runs without network access: valuations use the synthetic company,
the store and .env writes go to temporary directories, and the Anthropic SDK
talks to an in-process mock transport (no real API calls). Run with:
    python -m unittest tests.test_backend
"""

from __future__ import annotations

import dataclasses
import io
import json
import os
import shutil
import stat
import subprocess
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient

from backend import ai_service, exports, filings, store, valuation_service as vs
from backend import app as backend
from backend.serialization import build_ai_context
from equity_valuation import value_company
from equity_valuation.data.synthetic import (
    DEMO_PEERS,
    DEMO_TICKER,
    SyntheticProvider,
    make_company,
)
from equity_valuation.models.dcf import run_dcf
from equity_valuation.schemas import DCFAssumptions
from equity_valuation.utils import fade_path

LOCAL = "http://127.0.0.1:8000"


def _synthetic(ticker, refresh=False):
    return SyntheticProvider()


class _PricedProvider(SyntheticProvider):
    """The synthetic firm with a different market price and/or share count."""

    def __init__(self, price=None, share_mult=1.0):
        self.price, self.share_mult = price, share_mult

    def get_company_data(self, ticker):
        c = make_company()
        m = c.market
        price = self.price if self.price is not None else m.price / self.share_mult
        c.market = dataclasses.replace(
            m, price=price, shares_outstanding=m.shares_outstanding * self.share_mult)
        c.financials = dataclasses.replace(
            c.financials,
            diluted_shares=[s * self.share_mult for s in c.financials.diluted_shares])
        return c


def _report(provider, payload=None):
    macro, dcf_a, ddm_a, *_ = vs.parse_assumptions(payload or {})
    rep = value_company(DEMO_TICKER, provider=provider, macro=macro,
                        dcf_assumptions=dcf_a, ddm_assumptions=ddm_a, peers=DEMO_PEERS)
    return rep, macro, dcf_a


def _reprice(rep, macro, dcf_a, g1):
    path = vs._revenue_growth_path(g1, dcf_a.terminal_growth, dcf_a.forecast_years)
    a = dataclasses.replace(dcf_a, revenue_growth=path)
    return run_dcf(rep.company, macro, a, rep.current_price).implied_price


def _client():
    return TestClient(backend.app, base_url=LOCAL)


# --------------------------------------------------------------------------- #
#  Assumption parsing
# --------------------------------------------------------------------------- #
class ParseAssumptionTests(unittest.TestCase):
    def test_one_year_horizon_keeps_year1_growth(self):
        self.assertEqual(vs._revenue_growth_path(0.15, 0.025, 1), [0.15])
        for n in range(2, 16):
            self.assertEqual(vs._revenue_growth_path(0.15, 0.025, n),
                             fade_path(0.15, 0.025, n))
        _m, dcf_a, *_rest, echo = vs.parse_assumptions(
            {"forecast_years": 1, "revenue_growth_y1": 0.15})
        self.assertEqual(dcf_a.revenue_growth, [0.15])
        self.assertEqual(echo["revenue_growth"], [0.15])

    def test_non_finite_and_boolean_values_rejected(self):
        for key in ("rf", "erp", "tax_rate", "terminal_growth", "forecast_years",
                    "exit_ev_ebitda", "target_ebit_margin", "revenue_growth_y1"):
            for bad in ("inf", "-Infinity", "nan", float("inf"), float("nan"), 1e400, True):
                with self.subTest(key=key, value=bad):
                    with self.assertRaises(vs.AssumptionError) as caught:
                        vs.parse_assumptions({key: bad})
                    self.assertIn(key, str(caught.exception))
        with self.assertRaises(vs.AssumptionError):
            vs.parse_assumptions({"revenue_growth": ["x", None]})
        with self.assertRaises(vs.AssumptionError):
            vs.parse_assumptions({"revenue_growth": [0.1, float("nan")]})

    def test_lenient_inputs_still_parse(self):
        _m, dcf_a, _d, _p, toggles, echo = vs.parse_assumptions(
            {"forecast_years": "7", "rf": "0.04", "tax_rate": "", "revenue_growth": ["0.1", 0.05],
             "run_dcf": "false", "run_ddm": 0, "run_comps": "yes"})
        self.assertEqual(dcf_a.forecast_years, 7)
        self.assertEqual(echo["rf"], 0.04)
        self.assertIsNone(echo["tax_rate"])  # unparseable -> engine default
        self.assertEqual(dcf_a.revenue_growth, [0.1, 0.05])
        self.assertFalse(toggles["run_dcf"])
        self.assertFalse(toggles["run_ddm"])
        self.assertTrue(toggles["run_comps"])
        self.assertTrue(toggles["run_fcfe"])

    def test_api_returns_400_not_500(self):
        with patch.object(vs, "_get_provider", _synthetic):
            c = _client()
            for body in ('{"ticker":"SYNT","forecast_years":"inf"}',
                         '{"ticker":"SYNT","forecast_years":1e400}',
                         '{"ticker":"SYNT","terminal_growth":NaN}',
                         '{"ticker":"SYNT","rf":Infinity}'):
                with self.subTest(body=body):
                    r = c.post("/api/valuation", content=body,
                               headers={"content-type": "application/json"})
                    self.assertEqual(r.status_code, 400)
                    self.assertIsInstance(r.json()["detail"], str)
            r = c.post("/api/export/memo", json={"ticker": "SYNT", "rf": "nan"})
            self.assertEqual(r.status_code, 400)
            r = c.post("/api/valuation",
                       json={"ticker": "SYNT", "forecast_years": 1, "revenue_growth_y1": 0.15})
            self.assertEqual(r.status_code, 200)
            d = r.json()
            self.assertEqual(d["dcf"]["assumptions"]["revenue_growth_path"], [0.15])
            self.assertEqual(d["reverse_dcf"]["current_assumption_y1"], 0.15)


# --------------------------------------------------------------------------- #
#  Reverse DCF
# --------------------------------------------------------------------------- #
class _FakeDCF:
    """Stands in for run_dcf with a chosen price(g1) curve."""

    def __init__(self, fn):
        self.fn = fn

    def __call__(self, company, macro, a, price):
        return SimpleNamespace(implied_price=self.fn(a.revenue_growth[0]))


def _solve_fake(fn, price, n=5):
    rep = SimpleNamespace(dcf=object(), company=None, current_price=price)
    with patch("equity_valuation.models.dcf.run_dcf", _FakeDCF(fn)):
        return vs._reverse_dcf(rep, DCFAssumptions(forecast_years=n), None)


class ReverseDCFTests(unittest.TestCase):
    def test_solution_does_not_depend_on_share_count(self):
        results = []
        for mult in (1.0, 1_000.0, 100_000.0):
            rep, macro, dcf_a = _report(_PricedProvider(share_mult=mult))
            r = vs._reverse_dcf(rep, dcf_a, macro)
            self.assertTrue(r["converged"], mult)
            g1 = r["implied_growth_y1"]
            self.assertAlmostEqual(_reprice(rep, macro, dcf_a, g1) / rep.current_price, 1.0,
                                   delta=1e-6)
            results.append(g1)
        self.assertAlmostEqual(results[0], results[1], delta=1e-6)
        self.assertAlmostEqual(results[0], results[2], delta=1e-6)

    def test_penny_price_uses_relative_tolerance(self):
        r = _solve_fake(lambda g: 0.001 + 0.01 * g, price=0.001 + 0.01 * 0.1234)
        self.assertTrue(r["converged"])
        self.assertAlmostEqual(r["implied_growth_y1"], 0.1234, delta=1e-7)

    def test_root_inside_non_monotone_range_is_found(self):
        # Both ends of [-40%, 80%] price below the market, but the curve
        # crosses it twice in between; the old endpoint check said "outside".
        r = _solve_fake(lambda g: 50.0 - 400.0 * (g - 0.2) ** 2, price=40.0)
        self.assertTrue(r["converged"])
        self.assertAlmostEqual(r["implied_growth_y1"], 0.2 - (10 / 400) ** 0.5, delta=1e-7)
        self.assertIn("More than one", r["note"])

    def test_jump_across_price_is_not_reported_as_converged(self):
        r = _solve_fake(lambda g: 30.0 if g < 0.1 else 50.0, price=40.0)
        self.assertFalse(r["converged"])
        self.assertIsNone(r["implied_growth_y1"])

    def test_invalid_price_gives_no_solution(self):
        for price in (float("nan"), float("inf"), 0.0, -5.0, None):
            with self.subTest(price=price):
                r = _solve_fake(lambda g: 40.0 + g, price=price)
                self.assertFalse(r["converged"])
                self.assertIsNone(r["implied_growth_y1"])

    def test_one_year_horizon_solves_or_explains(self):
        payload = {"forecast_years": 1, "revenue_growth_y1": 0.15}
        # Market price inside the 1-year model's range: converges and reprices.
        rep, macro, dcf_a = _report(_PricedProvider(price=30.0), payload)
        r = vs._reverse_dcf(rep, dcf_a, macro)
        self.assertEqual(r["current_assumption_y1"], 0.15)
        self.assertTrue(r["converged"])
        self.assertAlmostEqual(_reprice(rep, macro, dcf_a, r["implied_growth_y1"]), 30.0,
                               delta=30.0 * 1e-6)
        # Default synthetic price is above what one year of growth can reach:
        # the note says so with the model's actual price range.
        rep, macro, dcf_a = _report(SyntheticProvider(), payload)
        r = vs._reverse_dcf(rep, dcf_a, macro)
        self.assertFalse(r["converged"])
        self.assertIn("only spans", r["note"])


# --------------------------------------------------------------------------- #
#  AI grounding context
# --------------------------------------------------------------------------- #
class AIContextTests(unittest.TestCase):
    def test_context_carries_the_drivers_actually_used(self):
        with patch.object(vs, "_get_provider", _synthetic):
            d = vs.run_valuation(DEMO_TICKER, {
                "revenue_growth_y1": 0.12, "target_ebit_margin": 0.3, "tax_rate": 0.23,
                "terminal_method": "exit_multiple", "exit_ev_ebitda": 14})
            default = vs.run_valuation(DEMO_TICKER, {})
        ctx = build_ai_context(d)
        self.assertIn("DCF revenue-growth path: 12.0%", ctx)
        self.assertIn("target 30.0%", ctx)
        self.assertIn("tax rate used 23.0%", ctx)
        self.assertIn("exit EV/EBITDA 14.0x", ctx)
        self.assertIn("revenue_growth_y1=0.12", ctx)
        self.assertIn("exit_ev_ebitda=14", ctx)
        ctx = build_ai_context(default)
        self.assertIn("DCF revenue-growth path: 8.0%", ctx)
        self.assertNotIn("tax n/a", ctx)
        self.assertIn("tax 21.0%", ctx)
        self.assertIn("REVERSE DCF", ctx)

    def test_sparse_report_does_not_raise(self):
        ctx = build_ai_context({"summary": {}, "dcf": None, "macro": {}})
        self.assertIn("CURRENT ASSUMPTION VALUES", ctx)


# --------------------------------------------------------------------------- #
#  DNS-rebinding guard
# --------------------------------------------------------------------------- #
class HostGuardTests(unittest.TestCase):
    def test_local_hosts_pass(self):
        c = TestClient(backend.app)
        for host in ("127.0.0.1:8000", "localhost:3000", "localhost", "LOCALHOST:8000",
                     "127.0.0.1", "[::1]:8000", "[::1]", "::1"):
            with self.subTest(host=host):
                self.assertEqual(c.get("/api/health", headers={"host": host}).status_code, 200)

    def test_next_rewrite_proxy_and_desktop_calls_pass(self):
        c = TestClient(backend.app)
        # Next's proxy (changeOrigin) sets Host to the backend target and
        # forwards the browser's Host as X-Forwarded-Host.
        for fwd in ("localhost:3000", "127.0.0.1:3000", ""):
            r = c.get("/api/health", headers={"host": "127.0.0.1:8000",
                                              "x-forwarded-host": fwd,
                                              "origin": "http://localhost:3000"})
            self.assertEqual(r.status_code, 200, fwd)
        # The desktop shell calls the backend's dynamic port directly.
        r = c.get("/api/health", headers={"host": "127.0.0.1:53817",
                                          "origin": "http://127.0.0.1:53816"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers.get("access-control-allow-origin"), "http://127.0.0.1:53816")

    def test_rebound_hosts_rejected(self):
        c = TestClient(backend.app)
        evil = "rebind.attacker.example:8000"
        with patch.object(backend, "_upsert_env_file") as save, \
                patch.dict(backend.os.environ, {}, clear=True):
            r = c.post("/api/settings", json={"anthropic_api_key": "sk-attacker"},
                       headers={"host": evil, "origin": f"http://{evil}"})
            self.assertEqual(r.status_code, 400)
            save.assert_not_called()
            self.assertNotIn("ANTHROPIC_API_KEY", backend.os.environ)
        for headers in ({"host": evil},
                        {"host": "localhost.evil.example"},
                        {"host": "127.0.0.1.evil.example:8000"},
                        {"host": "127.0.0.1:8000", "x-forwarded-host": evil},
                        {"host": "127.0.0.1:8000", "x-forwarded-host": f"localhost:3000, {evil}"}):
            with self.subTest(headers=headers):
                for path in ("/api/watchlist", "/api/research_state/SYNT", "/api/health"):
                    self.assertEqual(c.get(path, headers=headers).status_code, 400)


# --------------------------------------------------------------------------- #
#  .env persistence
# --------------------------------------------------------------------------- #
class EnvFileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.env = Path(self.tmp.name) / ".env"
        self.patch = patch.object(backend, "_ENV_PATH", self.env)
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    def test_new_file_is_owner_only(self):
        backend._upsert_env_file({"FMP_API_KEY": "fmp-new"})
        self.assertEqual(self.env.read_text(), "FMP_API_KEY=fmp-new\n")
        self.assertEqual(stat.S_IMODE(self.env.stat().st_mode), 0o600)
        self.assertEqual(os.listdir(self.tmp.name), [".env"])

    def test_every_assignment_replaced_once(self):
        self.env.write_text(
            "# keys\nFMP_API_KEY=old\nANTHROPIC_MODEL=claude-opus-5\n"
            "  export FMP_API_KEY = dupe\nFMP_API_KEY=dupe-later\nOTHER=1\n")
        os.chmod(self.env, 0o644)
        backend._upsert_env_file({"FMP_API_KEY": "fmp-new", "ANTHROPIC_API_KEY": "sk-1"})
        self.assertEqual(
            self.env.read_text(),
            "# keys\nFMP_API_KEY=fmp-new\nANTHROPIC_MODEL=claude-opus-5\nOTHER=1\n"
            "ANTHROPIC_API_KEY=sk-1\n")
        self.assertEqual(stat.S_IMODE(self.env.stat().st_mode), 0o600)
        bash = shutil.which("bash")
        if bash:
            out = subprocess.run(
                [bash, "-c", f"set -a; source '{self.env}'; echo \"$FMP_API_KEY\""],
                capture_output=True, text=True, check=True).stdout.strip()
            self.assertEqual(out, "fmp-new")

    def test_settings_route_writes_through_real_upsert(self):
        with patch.dict(backend.os.environ, {}, clear=True), patch.object(backend, "_fmp"):
            r = _client().post("/api/settings", json={"fmp_api_key": "fmp-abc"})
            self.assertEqual(r.status_code, 200)
            self.assertIn("FMP_API_KEY=fmp-abc\n", self.env.read_text())
            r = _client().post("/api/settings", json={"fmp_api_key": "x\nEVIL=1"})
            self.assertEqual(r.status_code, 400)
            self.assertNotIn("EVIL", self.env.read_text())


# --------------------------------------------------------------------------- #
#  Research store
# --------------------------------------------------------------------------- #
class StoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        data = Path(self.tmp.name) / "data"
        self.path = data / "copilot_store.json"
        self.patches = [patch.object(store, "_DATA_DIR", data),
                        patch.object(store, "_STORE_PATH", self.path)]
        for p in self.patches:
            p.start()
        data.mkdir()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        self.tmp.cleanup()

    def backups(self):
        return sorted(p for p in self.path.parent.iterdir() if ".corrupt" in p.name)

    def test_round_trip_leaves_no_temp_files(self):
        store.upsert_watchlist({"ticker": "aapl", "price": 1.0})
        store.save_research("AAPL", {"notes": "hello"})
        self.assertEqual([w["ticker"] for w in store.get_watchlist()], ["AAPL"])
        self.assertEqual(store.get_research("aapl")["notes"], "hello")
        self.assertEqual(os.listdir(self.path.parent), ["copilot_store.json"])

    def test_undecodable_store_is_backed_up_not_500(self):
        raw = b'{"watchlist": [], "research": {"X": {"notes": "caf\xe9"}}}'
        self.path.write_bytes(raw)
        r = _client().get("/api/watchlist")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json(), {"watchlist": []})
        (backup,) = self.backups()
        self.assertEqual(backup.read_bytes(), raw)

    def test_repeated_corruption_keeps_every_backup(self):
        for content in ("{bad1", "{bad2", '["a list, not a store"]', '{"watchlist": {}}',
                        '{"research": {"X": []}}'):
            self.path.write_text(content)
            self.assertEqual(store.get_watchlist(), [])
            store.upsert_watchlist({"ticker": "T"})
        self.assertEqual(sorted(b.read_text() for b in self.backups()),
                         sorted(['["a list, not a store"]', '{"research": {"X": []}}',
                                 '{"watchlist": {}}', "{bad1", "{bad2"]))

    def test_transient_read_error_does_not_move_valid_store(self):
        store.upsert_watchlist({"ticker": "KEEP"})
        before = self.path.read_bytes()
        with patch("backend.store.open", side_effect=PermissionError("locked"), create=True):
            with self.assertRaises(store.StoreError):
                store.get_watchlist()
            r = _client().get("/api/watchlist")
            self.assertEqual(r.status_code, 503)
            r = _client().post("/api/watchlist", json={"action": "add", "ticker": "NEW"})
            self.assertEqual(r.status_code, 503)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.backups(), [])
        self.assertEqual([w["ticker"] for w in store.get_watchlist()], ["KEEP"])


# --------------------------------------------------------------------------- #
#  Office exports
# --------------------------------------------------------------------------- #
class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(vs, "_get_provider", _synthetic),
                        patch.object(exports, "_ensure_out", lambda: Path(self.tmp.name))]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        self.tmp.cleanup()

    def test_number_formatting(self):
        self.assertEqual(exports._money(-41.17), "-$41.17")
        self.assertEqual(exports._money(12.5, "EUR"), "€12.50")
        self.assertEqual(exports._money(float("nan")), "n/a")
        self.assertEqual(exports._cap(3e8), "300M")
        self.assertEqual(exports._cap(3.4e9), "3.4B")
        self.assertEqual(exports._cap(2.345e12), "2,345B")
        self.assertEqual(exports._cap(None), "n/a")

    def test_memo_shows_exit_multiple_and_own_metadata(self):
        from docx import Document
        from pptx import Presentation

        payload = {"peers": ",".join(DEMO_PEERS), "terminal_method": "exit_multiple",
                   "exit_ev_ebitda": 14}
        memo = Document(exports.export_memo(DEMO_TICKER, payload))
        text = "\n".join(p.text for p in memo.paragraphs)
        self.assertIn("exit EV/EBITDA 14.0x", text)
        deck = Presentation(exports.export_deck(DEMO_TICKER, payload))
        year = __import__("datetime").datetime.now().year
        for props in (memo.core_properties, deck.core_properties):
            self.assertEqual(props.author, "Equity Research Automation")
            self.assertEqual(props.last_modified_by, "Equity Research Automation")
            self.assertEqual(props.comments, "")
            self.assertIn("Synthetic Corp", props.title)
            self.assertGreaterEqual(props.created.year, year - 1)

    def test_concurrent_same_ticker_exports_are_not_corrupted(self):
        from openpyxl import load_workbook

        c = _client()

        def run(g):
            r = c.post("/api/export/excel", json={"ticker": "SYNT", "revenue_growth_y1": g})
            self.assertEqual(r.status_code, 200)
            self.assertIn("SYNT_valuation.xlsx", r.headers["content-disposition"])
            load_workbook(io.BytesIO(r.content))  # raises on a torn file
            return True

        with ThreadPoolExecutor(3) as ex:
            self.assertTrue(all(ex.map(run, [0.01 * i for i in range(6)])))


# --------------------------------------------------------------------------- #
#  AI service (mock transport; never calls the real API)
# --------------------------------------------------------------------------- #
def _sdk_client(handler):
    import anthropic
    import httpx2

    return anthropic.Anthropic(
        api_key="sk-test", max_retries=0,
        http_client=anthropic.DefaultHttpxClient(transport=httpx2.MockTransport(handler)))


def _message(blocks, stop_reason="end_turn"):
    return {"id": "msg_1", "type": "message", "role": "assistant", "model": "claude-opus-5",
            "content": blocks, "stop_reason": stop_reason, "stop_sequence": None,
            "usage": {"input_tokens": 1, "output_tokens": 1}}


class _Recorder:
    def __init__(self, status=200, body=None, headers=None, exc=None):
        self.status, self.body, self.headers, self.exc = status, body, headers or {}, exc
        self.requests: list[dict] = []

    def __call__(self, request):
        import httpx2

        self.requests.append(json.loads(request.content))
        if self.exc is not None:
            raise self.exc
        return httpx2.Response(self.status, json=self.body, headers=self.headers)


class AIServiceTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {"ANTHROPIC_API_KEY": "sk-test"})
        self.env.start()

    def tearDown(self):
        self.env.stop()

    def _with(self, rec):
        return patch.object(ai_service, "_client", lambda: _sdk_client(rec))

    def test_chat_request_shape_and_reply(self):
        rec = _Recorder(body=_message([{"type": "thinking", "thinking": "", "signature": "s"},
                                       {"type": "text", "text": " The answer. "}]))
        with self._with(rec):
            reply = ai_service.chat("ctx", [{"role": "user", "content": "q"}])
        self.assertEqual(reply, "The answer.")
        req = rec.requests[0]
        self.assertGreaterEqual(req["max_tokens"], 16000)
        self.assertEqual(req["thinking"], {"type": "adaptive"})
        self.assertEqual(req["output_config"], {"effort": "high"})
        self.assertEqual(req["messages"], [{"role": "user", "content": "q"}])

    def test_chat_never_returns_empty(self):
        cases = [
            (_message([{"type": "thinking", "thinking": "", "signature": "s"}], "max_tokens"),
             "output limit"),
            (_message([], "refusal"), "declined"),
            (_message([{"type": "text", "text": "  "}]), "empty"),
        ]
        for body, words in cases:
            with self.subTest(stop=body["stop_reason"]), self._with(_Recorder(body=body)):
                with self.assertRaises(ai_service.AIError) as caught:
                    ai_service.chat("ctx", [{"role": "user", "content": "q"}])
                self.assertIn(words, str(caught.exception))
        body = _message([{"type": "text", "text": "partial"}], "max_tokens")
        with self._with(_Recorder(body=body)):
            reply = ai_service.chat("ctx", [{"role": "user", "content": "q"}])
        self.assertTrue(reply.startswith("partial"))
        self.assertIn("truncated", reply)

    def test_bad_turns_rejected_before_any_call(self):
        bad = [
            [],
            [{"content": "hi"}],
            [{"role": "system", "content": "hi"}],
            [{"role": "assistant", "content": "hi"}, {"role": "user", "content": "q"}],
            [{"role": "user", "content": "q"}, {"role": "assistant", "content": "prefill"}],
            [{"role": "user", "content": "q"}, {"role": "assistant", "content": ""},
             {"role": "user", "content": "q2"}],
            [{"role": "user", "content": ["block"]}],
            ["just a string"],
        ]
        rec = _Recorder(body=_message([{"type": "text", "text": "x"}]))
        with self._with(rec):
            for turns in bad:
                with self.subTest(turns=turns):
                    with self.assertRaises(ai_service.AIError):
                        ai_service.chat("ctx", turns)
            r = _client().post("/api/ai/chat", json={"turns": [{"content": "hi"}]})
            self.assertEqual(r.status_code, 400)
            self.assertIsInstance(r.json()["detail"], str)
        self.assertEqual(rec.requests, [])

    def test_pdf_payloads_are_normalised(self):
        blocks = ai_service._pdf_blocks([
            "not a dict", None, {"data_base64": ""},
            {"data_base64": "data:application/pdf;base64,JVBE\nRi0x\n", "name": "a.pdf"},
            {"data": " JVBE Ri0x "},
        ])
        self.assertEqual([b["source"]["data"] for b in blocks], ["JVBERi0x", "JVBERi0x"])
        self.assertEqual(blocks[0]["title"], "a.pdf")
        rec = _Recorder(body=_message([{"type": "text", "text": "ok"}]))
        with self._with(rec):
            ai_service.chat("ctx", [{"role": "user", "content": "q1"},
                                    {"role": "assistant", "content": "a1"},
                                    {"role": "user", "content": "q2"}],
                            pdfs=[{"data_base64": "JVBERi0x"}])
        last = rec.requests[0]["messages"][-1]
        self.assertEqual(last["role"], "user")
        self.assertEqual(last["content"][0]["type"], "document")
        self.assertEqual(last["content"][-1], {"type": "text", "text": "q2"})

    def test_haiku_override_omits_thinking_and_effort(self):
        digest = {"summary": "s", "sentiment": "neutral", "key_facts": [], "risks": [],
                  "catalysts": [], "suggested_assumptions": []}
        rec = _Recorder(body=_message([{"type": "text", "text": json.dumps(digest)}]))
        with self._with(rec), patch.object(ai_service, "MODEL", "claude-haiku-4-5"):
            self.assertEqual(ai_service.digest("ctx", "material"), digest)
            rec.body = _message([{"type": "text", "text": "hi"}])
            ai_service.chat("ctx", [{"role": "user", "content": "q"}])
        structured, chat = rec.requests
        self.assertNotIn("thinking", structured)
        self.assertEqual(set(structured["output_config"]), {"format"})
        self.assertNotIn("thinking", chat)
        self.assertNotIn("output_config", chat)
        self.assertEqual(chat["model"], "claude-haiku-4-5")
        rec.requests.clear()
        rec.body = _message([{"type": "text", "text": json.dumps(digest)}])
        with self._with(rec):  # default model keeps adaptive thinking + effort
            ai_service.digest("ctx", "material")
        self.assertEqual(rec.requests[0]["thinking"], {"type": "adaptive"})
        self.assertEqual(rec.requests[0]["output_config"]["effort"], "high")
        self.assertGreaterEqual(rec.requests[0]["max_tokens"], 16000)

    def test_sdk_errors_map_to_clean_http_errors(self):
        import anthropic
        import httpx2

        def err(status, message, headers=None):
            body = {"type": "error", "error": {"type": "x", "message": message}}
            return _Recorder(status=status, body=body, headers=headers)

        req = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
        cases = [
            (err(401, "invalid x-api-key"), 502, "ANTHROPIC_API_KEY"),
            (err(403, "no access"), 502, "ANTHROPIC_API_KEY"),
            (err(404, "model: nope"), 502, "ANTHROPIC_MODEL"),
            (err(400, "messages: bad thing"), 400, "messages: bad thing"),
            (err(413, "too big"), 413, "too big"),
            (err(429, "slow down", {"retry-after": "7"}), 429, "slow down"),
            (err(500, "boom"), 502, "boom"),
            (err(529, "overloaded"), 503, "overloaded"),
            (_Recorder(exc=httpx2.ConnectError("refused", request=req)), 503, "Connection"),
        ]
        body = {"turns": [{"role": "user", "content": "q"}]}
        for rec, code, words in cases:
            with self.subTest(code=code, words=words), self._with(rec):
                r = _client().post("/api/ai/chat", json=body)
                self.assertEqual(r.status_code, code)
                detail = r.json()["detail"]
                self.assertIsInstance(detail, str)
                self.assertIn(words, detail)
                self.assertNotIn("Error code:", detail)
                if code == 429:
                    self.assertEqual(r.headers.get("retry-after"), "7")
        self.assertTrue(issubclass(anthropic.RateLimitError, anthropic.APIStatusError))

    def test_installed_sdk_accepts_the_parameters_we_send(self):
        import inspect

        import anthropic

        params = inspect.signature(anthropic.resources.messages.Messages.create).parameters
        self.assertIn("output_config", params)
        self.assertIn("thinking", params)


# --------------------------------------------------------------------------- #
#  Filing section extraction
# --------------------------------------------------------------------------- #
class FilingSectionTests(unittest.TestCase):
    @staticmethod
    def _words(tag, n):
        return " ".join(f"{tag}{i}" for i in range(n))

    def test_cross_references_do_not_end_a_section(self):
        w = self._words
        text = "\n".join([
            "Item 1. Business  3", "Item 1A. Risk Factors  12",
            "Item 7. Management's Discussion and Analysis  30",
            "Item 1. Business",
            f"BUSINESS_START {w('b', 150)}. For risks, see Part I, Item 1A of this "
            f"Form 10-K under Risk Factors. {w('c', 150)} BUSINESS_END",
            "Item 1A. Risk Factors",
            f"RISK_START {w('r', 200)} RISK_END",
            "Item 1B. Unresolved Staff Comments", "None.",
            "Item 7. Management's Discussion and Analysis of Financial Condition",
            f"MDNA_START {w('m', 150)} read with the financial statements in Item 8. "
            f"Financial Statements and Supplementary Data. {w('n', 150)} Market risk is "
            f"in Item 7A, Quantitative and Qualitative Disclosures. {w('o', 150)} MDNA_END",
            "Item 7A. Quantitative and Qualitative Disclosures About Market Risk",
            w("q", 50),
            "Item 8. Financial Statements and Supplementary Data",
            w("f", 50),
        ])
        s = filings.extract_sections(text, "10-K")
        self.assertTrue(s["business"].rstrip().endswith("BUSINESS_END"))
        self.assertTrue(s["risk_factors"].rstrip().endswith("RISK_END"))
        self.assertTrue(s["mdna"].rstrip().endswith("MDNA_END"))

    def test_falls_back_to_inline_end_heading(self):
        w = self._words
        text = (f"Item 1A. Risk Factors\n{w('r', 200)} RISK_END Item 1B. Unresolved "
                f"Staff Comments {w('z', 200)}")
        s = filings.extract_sections(text, "10-K")
        self.assertTrue(s["risk_factors"].rstrip().endswith("RISK_END"))


if __name__ == "__main__":
    unittest.main()
