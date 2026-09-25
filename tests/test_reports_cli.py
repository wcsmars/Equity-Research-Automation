"""Regression tests for the Excel/HTML reports and the command-line interface.

The Excel checks evaluate the workbook's own live formulas (with the small
evaluator below, which covers the arithmetic, SUM and IF the exporter writes)
and compare them with the model's numbers, so they hold whatever the demo
figures are. Everything runs offline on the synthetic company. Run with:
python -m unittest tests.test_reports_cli
"""

from __future__ import annotations

import contextlib
import dataclasses
import io
import os
import re
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from equity_valuation import value_company
from equity_valuation.data.synthetic import (
    DEMO_PEERS,
    DEMO_TICKER,
    SyntheticProvider,
    make_company,
)
from equity_valuation.report.excel import write_excel
from equity_valuation.report.html import write_html
from equity_valuation.schemas import CompRow

MINORITY = 30e9
PREFERRED = 10e9


# --------------------------------------------------------------------------- #
#  Fixtures
# --------------------------------------------------------------------------- #
class _ClaimsProvider(SyntheticProvider):
    """The synthetic company with minority interest and preferred equity."""

    def get_company_data(self, ticker):
        company = make_company()
        company.balance_sheet.minority_interest = MINORITY
        company.balance_sheet.preferred_equity = PREFERRED
        return company


class _OutlierPeersProvider(SyntheticProvider):
    """Four peers, one with an outlying EV/EBITDA that the comps model trims."""

    def get_peer_comp_rows(self, tickers):
        return [
            CompRow(ticker=f"P{i}", name=f"P{i}", market_cap=5e10, enterprise_value=5.2e10,
                    ev_ebitda=m, ev_sales=4.0, pe=20.0, pb=5.0, peg=1.8)
            for i, m in enumerate((18.0, 22.0, 20.0, 90.0))
        ]


def _report(provider=None, **kwargs):
    kwargs.setdefault("peers", DEMO_PEERS)
    return value_company(DEMO_TICKER, provider=provider or SyntheticProvider(), **kwargs)


def _tmpdir(test: unittest.TestCase) -> str:
    tmp = tempfile.TemporaryDirectory()
    test.addCleanup(tmp.cleanup)
    return tmp.name


# --------------------------------------------------------------------------- #
#  A tiny evaluator for the exporter's formulas
# --------------------------------------------------------------------------- #
class _XlError(str):
    """An Excel error value such as #DIV/0!; propagates through arithmetic."""


_DIV0 = _XlError("#DIV/0!")
_VALUE = _XlError("#VALUE!")
_REF_RE = re.compile(r"(?:(\w+)!)?\$?([A-Z]{1,3})\$?(\d+)")
_TOKEN_RE = re.compile(
    r'\s*(?:(?P<str>"(?:[^"]|"")*")'
    r"|(?P<ref>(?:\w+!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![\w(])"
    r"|(?P<func>[A-Z]+)\("
    r"|(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)"
    r"|(?P<op><>|<=|>=|[-+*/=<>(),]))"
)


def _to_num(v):
    if isinstance(v, _XlError):
        return v
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return _VALUE  # text in arithmetic


def _arith(op, a, b):
    a, b = _to_num(a), _to_num(b)
    for x in (a, b):
        if isinstance(x, _XlError):
            return x
    if op == "+":
        return a + b
    if op == "-":
        return a - b
    if op == "*":
        return a * b
    return _DIV0 if b == 0 else a / b


def _compare(op, a, b):
    for x in (a, b):
        if isinstance(x, _XlError):
            return x
    if a is None:
        a = 0.0 if not isinstance(b, str) else ""
    if b is None:
        b = 0.0 if not isinstance(a, str) else ""
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
            "<=": a <= b, ">=": a >= b}[op]


class XlEval:
    """Evaluate a workbook's formulas (numbers, refs, ranges, + - * /, SUM, IF)."""

    def __init__(self, wb):
        self.wb = wb
        self._cache: dict = {}

    def value(self, sheet: str, coord: str):
        key = (sheet, coord)
        if key not in self._cache:
            raw = self.wb[sheet][coord].value
            if isinstance(raw, str) and raw.startswith("="):
                raw = self._formula(sheet, raw[1:])
            self._cache[key] = raw
        return self._cache[key]

    def formulas(self):
        """Yield (sheet, coord, evaluated value) for every formula cell."""
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        yield ws.title, c.coordinate, self.value(ws.title, c.coordinate)

    # -- recursive-descent parser producing lazy thunks ---------------------- #
    def _formula(self, sheet, text):
        toks, pos = [], 0
        text = text.strip()
        while pos < len(text):
            m = _TOKEN_RE.match(text, pos)
            if not m or m.end() == pos:
                raise ValueError(f"cannot parse formula {text!r} at {pos}")
            kind = m.lastgroup
            toks.append((kind, m.group(kind)))
            pos = m.end()
        self._toks, self._i, self._sheet = toks, 0, sheet
        node = self._comparison()
        if self._i != len(toks):
            raise ValueError(f"trailing tokens in {text!r}")
        return node()

    def _peek(self):
        return self._toks[self._i] if self._i < len(self._toks) else (None, None)

    def _take(self, value=None):
        tok = self._peek()
        if value is not None and tok[1] != value:
            raise ValueError(f"expected {value!r}, got {tok!r}")
        self._i += 1
        return tok

    def _comparison(self):
        left = self._additive()
        kind, op = self._peek()
        if kind == "op" and op in ("=", "<>", "<", ">", "<=", ">="):
            self._take()
            right = self._additive()
            return lambda: _compare(op, left(), right())
        return left

    def _additive(self):
        node = self._term()
        while self._peek()[1] in ("+", "-") and self._peek()[0] == "op":
            op = self._take()[1]
            rhs, lhs = self._term(), node
            node = (lambda o, a, b: lambda: _arith(o, a(), b()))(op, lhs, rhs)
        return node

    def _term(self):
        node = self._unary()
        while self._peek()[1] in ("*", "/") and self._peek()[0] == "op":
            op = self._take()[1]
            rhs, lhs = self._unary(), node
            node = (lambda o, a, b: lambda: _arith(o, a(), b()))(op, lhs, rhs)
        return node

    def _unary(self):
        if self._peek() == ("op", "-"):
            self._take()
            inner = self._unary()
            return lambda: _arith("-", 0.0, inner())
        if self._peek() == ("op", "+"):
            self._take()
            return self._unary()
        return self._primary()

    def _primary(self):
        kind, text = self._take()
        sheet = self._sheet
        if kind == "num":
            return lambda: float(text)
        if kind == "str":
            return lambda: text[1:-1].replace('""', '"')
        if kind == "ref":
            if ":" in text:
                return self._range(sheet, text)
            m = _REF_RE.fullmatch(text)
            ref_sheet = m.group(1) or sheet
            return lambda: self.value(ref_sheet, m.group(2) + m.group(3))
        if kind == "op" and text == "(":
            node = self._comparison()
            self._take(")")
            return node
        if kind == "func":
            args = []
            if self._peek() != ("op", ")"):
                args.append(self._comparison())
                while self._peek() == ("op", ","):
                    self._take()
                    args.append(self._comparison())
            self._take(")")
            return self._call(text, args)
        raise ValueError(f"unexpected token {text!r}")

    def _range(self, sheet, text):
        from openpyxl.utils import column_index_from_string, get_column_letter

        a, b = text.split(":")
        ma, mb = _REF_RE.fullmatch(a), _REF_RE.fullmatch(b)
        ref_sheet = ma.group(1) or sheet
        c1, c2 = column_index_from_string(ma.group(2)), column_index_from_string(mb.group(2))
        r1, r2 = int(ma.group(3)), int(mb.group(3))
        coords = [f"{get_column_letter(c)}{r}"
                  for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return lambda: [self.value(ref_sheet, c) for c in coords]

    def _call(self, name, args):
        if name == "SUM":
            def _sum():
                total = 0.0
                for arg in args:
                    vals = arg()
                    for v in (vals if isinstance(vals, list) else [vals]):
                        if isinstance(v, _XlError):
                            return v
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            total += v
                return total
            return _sum
        if name == "IF":
            def _if():
                cond = args[0]()
                if isinstance(cond, _XlError):
                    return cond
                if cond:
                    return args[1]()
                return args[2]() if len(args) > 2 else False
            return _if
        raise ValueError(f"unsupported function {name}")


def _label_rows(ws) -> dict:
    """{column-A label: row number} for a sheet (first occurrence wins)."""
    rows = {}
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(cell.value, str):
            rows.setdefault(cell.value.strip(), cell.row)
    return rows


def _workbook(test, report):
    path = write_excel(report, os.path.join(_tmpdir(test), "SYNT_valuation.xlsx"))
    wb = load_workbook(path)
    return wb, XlEval(wb)


# --------------------------------------------------------------------------- #
#  Excel
# --------------------------------------------------------------------------- #
class ExcelReconciliationTests(unittest.TestCase):
    def assertClose(self, got, want, msg=None):
        self.assertIsInstance(got, float, msg)
        self.assertAlmostEqual(got, want, delta=max(1e-9 * abs(want), 1e-9), msg=msg)

    def test_dcf_bridge_subtracts_minority_interest_and_preferred(self):
        report = _report(_ClaimsProvider())
        dcf = report.dcf
        self.assertIsNotNone(dcf)
        wb, ev = _workbook(self, report)
        rows = _label_rows(wb["DCF"])
        self.assertEqual(wb["DCF"][f"B{rows['Less: minority interest']}"].value, MINORITY)
        self.assertEqual(wb["DCF"][f"B{rows['Less: preferred equity']}"].value, PREFERRED)

        def cell(label):
            return ev.value("DCF", f"B{rows[label]}")

        self.assertClose(cell("Enterprise value"), dcf.enterprise_value)
        self.assertClose(cell("Equity value"), dcf.equity_value)
        self.assertClose(cell("Implied price / share"), dcf.implied_price)
        self.assertClose(cell("Upside / (downside)"), dcf.upside)
        # The claims actually move the price, so this is not a vacuous check.
        self.assertAlmostEqual(
            dcf.enterprise_value - dcf.net_debt - dcf.equity_value, MINORITY + PREFERRED,
            delta=1.0,
        )

    def test_demo_formulas_reconcile_to_the_model(self):
        report = _report()
        wb, ev = _workbook(self, report)

        rows = _label_rows(wb["DCF"])
        self.assertClose(ev.value("DCF", f"B{rows['Equity value']}"), report.dcf.equity_value)
        self.assertClose(ev.value("DCF", f"B{rows['Implied price / share']}"),
                         report.dcf.implied_price)

        rows = _label_rows(wb["DDM_FCFE"])
        self.assertClose(ev.value("DDM_FCFE", f"B{rows['Equity value']}"),
                         report.fcfe.equity_value)
        self.assertClose(ev.value("DDM_FCFE", f"B{rows['Implied price / share']}"),
                         report.fcfe.implied_price)

        rows = _label_rows(wb["Summary"])
        cur = report.current_price
        blended = report.summary["blended_target"]
        self.assertClose(wb["Summary"][f"B{rows['Blended target (median)']}"].value, blended)
        self.assertClose(ev.value("Summary", f"C{rows['Blended target (median)']}"),
                         blended / cur - 1.0)
        self.assertClose(ev.value("Summary", f"C{rows['DCF (FCFF)']}"),
                         report.dcf.implied_price / cur - 1.0)

        errors = [(s, c, v) for s, c, v in ev.formulas() if isinstance(v, _XlError)]
        self.assertEqual(errors, [])

    def test_zero_revenue_and_zero_price_give_no_div0(self):
        report = _report()
        n = len(report.dcf.years)
        report.dcf = dataclasses.replace(
            report.dcf, revenue=[0.0] * n, ebit=[0.0] * n, nopat=[0.0] * n,
        )
        report.current_price = 0.0
        _wb, ev = _workbook(self, report)
        errors = [(s, c, v) for s, c, v in ev.formulas() if isinstance(v, _XlError)]
        self.assertEqual(errors, [])

    def test_blended_target_stays_blank_when_engine_has_none(self):
        report = _report(run_dcf=False, run_comps=False, run_ddm=False, run_fcfe=False)
        self.assertIsNone(report.summary.get("blended_target"))
        wb, _ev = _workbook(self, report)
        rows = _label_rows(wb["Summary"])
        self.assertIsNone(wb["Summary"][f"B{rows['Blended target (median)']}"].value)

    def test_ddm_detail_formats_and_lists(self):
        report = _report()
        self.assertIsNotNone(report.ddm)
        wb, _ev = _workbook(self, report)
        ws = wb["DDM_FCFE"]
        rows = _label_rows(ws)
        self.assertIn("%", ws[f"B{rows['cost_of_equity']}"].number_format)
        if "high_growth_years" in rows:
            fmt = ws[f"B{rows['high_growth_years']}"].number_format
            self.assertNotIn("%", fmt)
            self.assertNotIn("$", fmt)
        for (cell,) in ws.iter_rows(min_col=2, max_col=2):
            if isinstance(cell.value, str):
                self.assertFalse(cell.value.startswith("["), f"{cell.coordinate}: {cell.value}")
        dividends = report.ddm.detail.get("dividends")
        if dividends:
            r = rows["dividends"]
            spilled = [ws.cell(row=r, column=2 + j).value for j in range(len(dividends))]
            self.assertEqual(spilled, [float(d) for d in dividends])

    def test_projection_headers_label_the_same_fiscal_years(self):
        report = _report()
        wb, _ev = _workbook(self, report)

        def header(ws, first_label):
            r = _label_rows(ws)[first_label]
            return [c.value for c in ws[r][1:] if c.value]

        dcf_hdr = header(wb["DCF"], "(values in reporting currency)")
        fcfe_hdr = header(wb["DDM_FCFE"], "(reporting currency)")
        self.assertEqual(len(dcf_hdr), len(report.dcf.years))
        self.assertEqual(dcf_hdr, fcfe_hdr[: len(dcf_hdr)])
        self.assertNotIn("FY 1", dcf_hdr)

    def test_control_characters_do_not_sink_the_workbook(self):
        report = _report()
        report.warnings.append("Comps: peer fetch failed: bad byte \x1b[0m in response")
        wb, _ev = _workbook(self, report)
        texts = [c.value for row in wb["Summary"].iter_rows() for c in row
                 if isinstance(c.value, str)]
        self.assertTrue(any("bad byte [0m in response" in t for t in texts))

    def test_one_failing_sheet_does_not_sink_the_workbook(self):
        from equity_valuation.report import excel

        report = _report()
        with patch.object(excel, "_write_comps", side_effect=RuntimeError("boom")):
            wb, _ev = _workbook(self, report)
        self.assertEqual(wb.sheetnames, ["Summary", "DCF", "Comps", "DDM_FCFE", "Sensitivity"])
        texts = [c.value for row in wb["Comps"].iter_rows() for c in row if c.value]
        self.assertTrue(any("boom" in str(t) for t in texts))


# --------------------------------------------------------------------------- #
#  HTML
# --------------------------------------------------------------------------- #
def _html(test, report, name="SYNT_valuation.html") -> str:
    path = write_html(report, os.path.join(_tmpdir(test), name))
    return Path(path).read_text(encoding="utf-8")


def _plotly_copies(text: str) -> int:
    return text.count("window.PlotlyConfig")


def _card(text: str, label: str) -> str:
    m = re.search(r"card-label'>" + re.escape(label) + r"</div><div class='card-value "
                  r"\w+'>([^<]*)<", text)
    return m.group(1) if m else None


class HtmlReportTests(unittest.TestCase):
    def test_bridge_shows_minority_interest_and_preferred(self):
        from equity_valuation.report.html import _fmt_big

        report = _report(_ClaimsProvider())
        text = _html(self, report)
        self.assertIn(f"Less: minority interest</th><td>{_fmt_big(MINORITY, '$')}<", text)
        self.assertIn(f"Less: preferred equity</th><td>{_fmt_big(PREFERRED, '$')}<", text)
        self.assertIn(f"Equity value</th><td>{_fmt_big(report.dcf.equity_value, '$')}<", text)

    def test_each_render_embeds_plotly_exactly_once_even_when_nested(self):
        # A second render starting while the first is in progress used to reset
        # the shared module state, leaving the first document with no plotly.js.
        from equity_valuation.report import html as html_mod

        report = _report()
        out = _tmpdir(self)
        original = html_mod._header_html
        calls = []

        def nested(*args, **kwargs):
            if not calls:
                calls.append(1)
                write_html(report, os.path.join(out, "inner.html"))
            return original(*args, **kwargs)

        with patch.object(html_mod, "_header_html", nested):
            write_html(report, os.path.join(out, "outer.html"))
        for name in ("outer.html", "inner.html"):
            text = Path(out, name).read_text(encoding="utf-8")
            self.assertEqual(_plotly_copies(text), 1, name)

    def test_concurrent_renders_embed_plotly_exactly_once(self):
        report = _report()
        out = _tmpdir(self)
        errors = []

        def work(i):
            try:
                for j in range(2):
                    write_html(report, os.path.join(out, f"r{i}_{j}.html"))
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=work, args=(i,)) for i in range(3)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(errors, [])
        for name in sorted(os.listdir(out)):
            text = Path(out, name).read_text(encoding="utf-8")
            self.assertEqual(_plotly_copies(text), 1, name)

    def test_no_blended_target_is_invented(self):
        report = _report(run_dcf=False, run_comps=False, run_ddm=False, run_fcfe=False)
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"), "n/a")
        self.assertEqual(_card(text, "Upside / downside"), "n/a")

    def test_header_lists_every_method_behind_the_blended_target(self):
        from equity_valuation.report.html import _fmt_price

        report = _report()
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"),
                         _fmt_price(report.summary["blended_target"], "$"))
        self.assertIn("Valuation by method", text)
        for name, price in report.summary["methods"].items():
            self.assertIn(f"{name}</th><td>{_fmt_price(price, '$')}<", text)

    def test_comps_chart_median_matches_the_trimmed_stats(self):
        from equity_valuation.report.html import _fmt_mult

        report = _report(_OutlierPeersProvider(), peers=["P0", "P1", "P2", "P3"])
        med = report.comps.stats["ev_ebitda"]["median"]
        text = _html(self, report)
        self.assertIn(f"Peer median {_fmt_mult(med)}", text)


# --------------------------------------------------------------------------- #
#  CLI
# --------------------------------------------------------------------------- #
def _run_cli(argv):
    """(exit code, stdout, stderr) of cli.main, with SystemExit mapped to its code."""
    from equity_valuation.cli import main

    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            code = main(argv)
        except SystemExit as exc:
            code = exc.code
    return code, out.getvalue(), err.getvalue()


class CliTests(unittest.TestCase):
    def test_terminal_growth_reaches_ddm_and_fcfe(self):
        from equity_valuation import engine

        seen = {}
        real = engine.value_company

        def spy(*args, **kwargs):
            seen.update(kwargs)
            return real(*args, **kwargs)

        with patch.object(engine, "value_company", spy):
            code, _out, err = _run_cli(["--demo", "--terminal-growth", "0.04", "--quiet",
                                        "--excel", "--out", _tmpdir(self)])
        self.assertEqual(code, 0, err)
        self.assertEqual(seen["dcf_assumptions"].terminal_growth, 0.04)
        self.assertEqual(seen["ddm_assumptions"].terminal_growth, 0.04)

    def test_bad_numeric_flags_are_rejected(self):
        cases = [
            (["--rf", "4.3"], "0.043"),
            (["--tax", "21"], "0.21"),
            (["--target-ebit-margin", "25"], "decimals"),
            (["--rf", "nan"], "finite"),
            (["--erp", "inf"], "finite"),
            (["--terminal-growth", "nan"], "finite"),
            (["--cost-of-debt", "-0.01"], ">= 0"),
            (["--forecast-years", "0"], "between 1 and"),
            (["--forecast-years", "-3"], "between 1 and"),
            (["--exit-ev-ebitda", "-5", "--terminal-method", "exit_multiple"], "> 0"),
            (["--terminal-method", "exit_multiple"], "requires --exit-ev-ebitda"),
        ]
        out = _tmpdir(self)
        for extra, hint in cases:
            with self.subTest(args=extra):
                code, _out, err = _run_cli(["--demo", "--quiet", "--out", out] + extra)
                self.assertEqual(code, 2)
                self.assertIn(hint, err)
                self.assertNotIn("Traceback", err)
        self.assertEqual(os.listdir(out), [])

    def test_valid_decimal_flags_still_run(self):
        code, _out, err = _run_cli([
            "--demo", "--quiet", "--excel", "--out", _tmpdir(self), "--rf", "-0.005",
            "--tax", "0.21", "--terminal-method", "exit_multiple", "--exit-ev-ebitda", "12",
        ])
        self.assertEqual(code, 0, err)

    def test_export_failure_exits_nonzero(self):
        out = _tmpdir(self)
        os.mkdir(os.path.join(out, "SYNT_valuation.xlsx"))
        os.mkdir(os.path.join(out, "SYNT_valuation.html"))
        code, _out, err = _run_cli(["--demo", "--quiet", "--out", out])
        self.assertEqual(code, 1)
        self.assertIn("Excel export failed", err)
        self.assertIn("HTML export failed", err)

    def test_out_pointing_at_a_file_is_a_clean_error(self):
        path = os.path.join(_tmpdir(self), "not_a_dir")
        Path(path).write_text("x", encoding="utf-8")
        code, _out, err = _run_cli(["--demo", "--quiet", "--out", path])
        self.assertEqual(code, 2)
        self.assertIn("not a directory", err)
        self.assertNotIn("Traceback", err)

    def test_ticker_is_sanitized_in_output_filenames(self):
        from equity_valuation import engine

        report = _report()
        report.company = dataclasses.replace(report.company, ticker="BRK/B")
        out = _tmpdir(self)
        with patch.object(engine, "value_company", lambda *a, **k: report):
            code, _stdout, err = _run_cli(["BRK/B", "--quiet", "--excel", "--out", out])
        self.assertEqual(code, 0, err)
        self.assertEqual(os.listdir(out), ["BRK_B_valuation.xlsx"])

    def test_demo_rejects_real_peer_tickers(self):
        code, _out, err = _run_cli(["--demo", "--peers", "MSFT,GOOGL", "--quiet",
                                    "--out", _tmpdir(self)])
        self.assertEqual(code, 2)
        self.assertIn("synthetic peers", err)

        code, _out, err = _run_cli(["--demo", "--peers", "peer1,PEER3", "--quiet", "--excel",
                                    "--out", _tmpdir(self)])
        self.assertEqual(code, 0, err)


if __name__ == "__main__":
    unittest.main()
