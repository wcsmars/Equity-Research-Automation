"""Offline checks for the demo run, every export format, and the reverse DCF.

All tests use the synthetic company from ``equity_valuation.data.synthetic``,
so none of them touch the network. Run with:  python -m unittest tests.test_exports
"""

from __future__ import annotations

import contextlib
import dataclasses
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from equity_valuation import value_company
from equity_valuation.data.synthetic import DEMO_PEERS, DEMO_TICKER, SyntheticProvider


def _demo_report():
    return value_company(DEMO_TICKER, provider=SyntheticProvider(), peers=DEMO_PEERS)


class DemoResultTests(unittest.TestCase):
    def test_summary_matches_readme_results(self):
        # The README "Results" table quotes these figures; keep them in step.
        s = _demo_report().summary
        self.assertAlmostEqual(s["current_price"], 40.84, delta=0.005)
        expected = {"DCF": 33.38, "Comps (median)": 42.88, "DDM": 10.99, "FCFE": 31.25}
        self.assertEqual(set(s["methods"]), set(expected))
        for method, price in expected.items():
            self.assertAlmostEqual(s["methods"][method], price, delta=0.005, msg=method)
        self.assertAlmostEqual(s["blended_target"], 32.32, delta=0.005)
        self.assertAlmostEqual(s["blended_upside"], -0.209, delta=0.0005)
        self.assertEqual(s["recommendation"], "Overvalued")

    def test_cli_demo_writes_excel_and_html(self):
        from equity_valuation.cli import main

        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(["--demo", "--quiet", "--out", tmp]), 0)
            self.assertTrue(os.path.isfile(os.path.join(tmp, "SYNT_valuation.xlsx")))
            self.assertTrue(os.path.isfile(os.path.join(tmp, "SYNT_valuation.html")))

    def test_cli_demo_rejects_a_live_ticker(self):
        from equity_valuation.cli import main

        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            main(["AAPL", "--demo"])


class EngineExportTests(unittest.TestCase):
    def test_excel_keeps_formulas_and_html_renders(self):
        from openpyxl import load_workbook

        from equity_valuation.report.excel import write_excel
        from equity_valuation.report.html import write_html

        report = _demo_report()
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = write_excel(report, os.path.join(tmp, "SYNT_valuation.xlsx"))
            html = write_html(report, os.path.join(tmp, "SYNT_valuation.html"))

            wb = load_workbook(xlsx)
            for sheet in ("Summary", "DCF", "Comps", "DDM_FCFE", "Sensitivity"):
                self.assertIn(sheet, wb.sheetnames)
            formulas = [
                c for row in wb["DCF"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")
            ]
            self.assertTrue(formulas, "DCF sheet should contain live formulas")

            text = Path(html).read_text(encoding="utf-8")
            self.assertIn("Synthetic Corp", text)
            self.assertIn("<html", text.lower())


class OfficeExportTests(unittest.TestCase):
    """The dashboard's Word memo and PowerPoint briefing, without a live provider."""

    def test_memo_and_deck(self):
        from docx import Document
        from pptx import Presentation

        from backend import exports, valuation_service

        payload = {"peers": ",".join(DEMO_PEERS)}
        with tempfile.TemporaryDirectory() as tmp, patch.object(
            valuation_service, "_get_provider",
            lambda ticker, refresh=False: SyntheticProvider(),
        ), patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            memo = exports.export_memo(DEMO_TICKER, payload)
            deck = exports.export_deck(DEMO_TICKER, payload)

            self.assertEqual(Path(memo).parent, Path(tmp))
            text = "\n".join(p.text for p in Document(memo).paragraphs)
            self.assertIn("Synthetic Corp", text)
            self.assertIn("Overvalued", text)
            self.assertGreaterEqual(len(Presentation(deck).slides), 1)


class ReverseDCFTests(unittest.TestCase):
    def test_solved_growth_reprices_to_market(self):
        from backend import valuation_service as vs
        from equity_valuation.models.dcf import run_dcf

        macro, dcf_a, ddm_a, _peers, _toggles, _echo = vs.parse_assumptions({})
        report = value_company(
            DEMO_TICKER, provider=SyntheticProvider(), macro=macro,
            dcf_assumptions=dcf_a, ddm_assumptions=ddm_a, peers=DEMO_PEERS,
        )
        result = vs._reverse_dcf(report, dcf_a, macro)
        self.assertTrue(result["converged"])
        g1 = result["implied_growth_y1"]
        # The synthetic firm grew 8% a year; a 20x P/E price needs more than that.
        self.assertGreater(g1, 0.08)

        path = vs._revenue_growth_path(g1, dcf_a.terminal_growth, dcf_a.forecast_years)
        solved = dataclasses.replace(dcf_a, revenue_growth=path)
        price = run_dcf(report.company, macro, solved, report.current_price).implied_price
        self.assertAlmostEqual(price, report.current_price, delta=0.01)


if __name__ == "__main__":
    unittest.main()
