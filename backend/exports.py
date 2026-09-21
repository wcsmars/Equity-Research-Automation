"""Office artifact generation — the analyst-facing outputs.

  * Excel model  -> engine's write_excel (live formulas, multi-tab)   [.xlsx]
  * HTML report  -> engine's write_html (interactive, offline)        [.html]
  * Research memo-> python-docx; model summary + the AI research note [.docx]
  * Briefing deck-> python-pptx; valuation summary, football field,
                    comps, thesis/risks                               [.pptx]

The memo/deck render fine with NO AI note (model-only); when a generated
research note dict is supplied they include the full cited write-up.
"""

from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path
from typing import Optional

from .valuation_service import run_valuation_report

_OUT_DIR = Path(__file__).resolve().parent.parent / "output"

_SYM = {"USD": "$", "EUR": "€", "GBP": "£", "JPY": "¥"}


def _money(x, cur="USD", dec=2) -> str:
    if x is None:
        return "n/a"
    return f"{_SYM.get(cur, '')}{x:,.{dec}f}"


def _pct(x, signed=False) -> str:
    if x is None:
        return "n/a"
    s = "+" if (signed and x > 0) else ""
    return f"{s}{x * 100:.1f}%"


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "_", name)


def _ensure_out() -> Path:
    os.makedirs(_OUT_DIR, exist_ok=True)
    return _OUT_DIR


# --------------------------------------------------------------------------- #
#  Engine-native exports
# --------------------------------------------------------------------------- #
def export_excel(ticker: str, payload: dict) -> str:
    from equity_valuation.report.excel import write_excel

    report, _ = run_valuation_report(ticker, payload)
    path = _ensure_out() / f"{_safe(ticker.upper())}_valuation.xlsx"
    return write_excel(report, str(path))


def export_html(ticker: str, payload: dict) -> str:
    from equity_valuation.report.html import write_html

    report, _ = run_valuation_report(ticker, payload)
    path = _ensure_out() / f"{_safe(ticker.upper())}_valuation.html"
    return write_html(report, str(path))


# --------------------------------------------------------------------------- #
#  Research memo (.docx)
# --------------------------------------------------------------------------- #
def export_memo(ticker: str, payload: dict, note: Optional[dict] = None) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor

    report, echo = run_valuation_report(ticker, payload)
    s = report.summary
    cur = s.get("currency", "USD")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = note.get("title") if note else None
    doc.add_heading(
        title or f"{s.get('name')} ({s.get('ticker')}) — Research Memo", level=0
    )
    meta = doc.add_paragraph()
    run = meta.add_run(
        f"Prepared {date.today().isoformat()} · price {_money(s.get('current_price'), cur)} · "
        f"blended fair value {_money(s.get('blended_target'), cur)} "
        f"({_pct(s.get('blended_upside'), signed=True)}) · model verdict: {s.get('recommendation')}"
    )
    run.font.color.rgb = RGBColor(0x60, 0x70, 0x8A)

    if note and note.get("stance"):
        doc.add_paragraph(f"Analytical stance: {note['stance']}").runs[0].bold = True

    # --- valuation summary table -------------------------------------------- #
    doc.add_heading("Valuation summary", level=1)
    methods = s.get("methods") or {}
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Method", "Implied value", "Vs. price"
    price = s.get("current_price")
    for name, val in methods.items():
        row = table.add_row().cells
        row[0].text = name
        row[1].text = _money(val, cur)
        row[2].text = _pct((val / price - 1) if (val and price) else None, signed=True)
    row = table.add_row().cells
    row[0].text = "Blended target"
    row[1].text = _money(s.get("blended_target"), cur)
    row[2].text = _pct(s.get("blended_upside"), signed=True)

    # --- key assumptions ------------------------------------------------------ #
    doc.add_heading("Key model assumptions", level=1)
    wacc = report.dcf.wacc.wacc if report.dcf else None
    bullets = [
        f"WACC: {_pct(wacc)} · risk-free {_pct(echo.get('rf'))} · ERP {_pct(echo.get('erp'))}",
        f"Forecast horizon: {echo.get('forecast_years')} years · terminal growth {_pct(echo.get('terminal_growth'))} ({echo.get('terminal_method')})",
    ]
    if echo.get("revenue_growth_y1") is not None:
        bullets.append(f"Year-1 revenue growth: {_pct(echo.get('revenue_growth_y1'))} (fading to terminal)")
    if echo.get("target_ebit_margin") is not None:
        bullets.append(f"Terminal EBIT margin: {_pct(echo.get('target_ebit_margin'))}")
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    # --- AI research note sections -------------------------------------------- #
    def _section(heading: str, items, bullet=True):
        if not items:
            return
        doc.add_heading(heading, level=1)
        if isinstance(items, str):
            doc.add_paragraph(items)
        else:
            for it in items:
                doc.add_paragraph(str(it), style="List Bullet" if bullet else None)

    if note:
        _section("Executive summary", note.get("executive_summary"))
        _section("Thesis", note.get("thesis"))
        _section("Valuation view", note.get("valuation_view"))
        _section("Key drivers", note.get("key_drivers"))
        _section("Risks", note.get("risks"))
        _section("Red flags / diligence items", note.get("red_flags"))
        _section("Catalysts", note.get("catalysts"))
        _section("What would change my mind", note.get("what_would_change_my_mind"))
        cites = note.get("citations") or []
        if cites:
            doc.add_heading("Sources", level=1)
            for c in cites:
                doc.add_paragraph(
                    f"{c.get('source')}: {c.get('note')}", style="List Bullet"
                )

    # --- footnotes -------------------------------------------------------------- #
    if report.warnings:
        doc.add_heading("Model notes", level=1)
        for w in report.warnings:
            doc.add_paragraph(str(w), style="List Bullet")
    p = doc.add_paragraph()
    r = p.add_run(
        "Generated by Equity Research Automation. Model estimates driven by user "
        "assumptions — research support, not investment advice."
    )
    r.italic = True
    r.font.size = Pt(8.5)

    path = _ensure_out() / f"{_safe(ticker.upper())}_research_memo.docx"
    doc.save(str(path))
    return str(path)


# --------------------------------------------------------------------------- #
#  Briefing deck (.pptx)
# --------------------------------------------------------------------------- #
def export_deck(ticker: str, payload: dict, note: Optional[dict] = None) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    report, _echo = run_valuation_report(ticker, payload)
    s = report.summary
    cur = s.get("currency", "USD")
    price = s.get("current_price")

    INK = RGBColor(0x12, 0x18, 0x26)
    DIM = RGBColor(0x5C, 0x6B, 0x84)
    BLUE = RGBColor(0x2F, 0x6B, 0xD8)
    GREEN = RGBColor(0x1F, 0x9D, 0x6F)
    ROSE = RGBColor(0xD6, 0x45, 0x61)
    AMBER = RGBColor(0xC8, 0x8A, 0x1A)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_slide():
        return prs.slides.add_slide(blank)

    def text(slide, l, t, w, h, s_, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = s_
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return box

    def bullets(slide, l, t, w, h, items, size=13):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for it in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = f"•  {it}"
            r.font.size = Pt(size)
            r.font.color.rgb = INK

    # --- slide 1: title --------------------------------------------------------- #
    sl = add_slide()
    text(sl, 0.7, 2.2, 11.9, 1.0, f"{s.get('name')}  ({s.get('ticker')})", 40, True)
    text(sl, 0.7, 3.3, 11.9, 0.6,
         f"Valuation briefing · {date.today().isoformat()}", 18, False, DIM)
    verdict = s.get("recommendation") or "n/a"
    vcolor = GREEN if verdict == "Undervalued" else ROSE if verdict == "Overvalued" else AMBER
    text(sl, 0.7, 4.2, 11.9, 0.6,
         f"Price {_money(price, cur)} · blended fair value {_money(s.get('blended_target'), cur)} "
         f"({_pct(s.get('blended_upside'), signed=True)}) · {verdict}", 18, True, vcolor)

    # --- slide 2: method summary ------------------------------------------------- #
    sl = add_slide()
    text(sl, 0.7, 0.4, 12, 0.6, "Valuation summary", 24, True)
    methods = list((s.get("methods") or {}).items())
    rows = len(methods) + 2
    tbl = sl.shapes.add_table(rows, 3, Inches(0.7), Inches(1.2), Inches(8.0),
                              Inches(0.4 * rows)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(0, 2).text = (
        "Method", "Implied value", "Vs. price")
    for i, (name, val) in enumerate(methods, start=1):
        tbl.cell(i, 0).text = name
        tbl.cell(i, 1).text = _money(val, cur)
        tbl.cell(i, 2).text = _pct((val / price - 1) if (val and price) else None, signed=True)
    tbl.cell(rows - 1, 0).text = "Blended target"
    tbl.cell(rows - 1, 1).text = _money(s.get("blended_target"), cur)
    tbl.cell(rows - 1, 2).text = _pct(s.get("blended_upside"), signed=True)

    # --- slide 3: football field --------------------------------------------------- #
    rows_ff = report.football_field or []
    if rows_ff:
        sl = add_slide()
        text(sl, 0.7, 0.4, 12, 0.6, "Football field — value ranges by method", 24, True)
        lows = [r.low for r in rows_ff] + [price]
        highs = [r.high for r in rows_ff] + [price]
        vmin, vmax = min(lows), max(highs)
        pad = (vmax - vmin) * 0.06 or 1.0
        vmin -= pad
        vmax += pad
        span = vmax - vmin
        chart_l, chart_w = 2.8, 9.4
        y = 1.4
        for r in rows_ff:
            text(sl, 0.6, y - 0.06, 2.1, 0.4, r.method, 12, False, DIM)
            x0 = chart_l + (r.low - vmin) / span * chart_w
            x1 = chart_l + (r.high - vmin) / span * chart_w
            bar = sl.shapes.add_shape(1, Inches(x0), Inches(y),
                                      Inches(max(x1 - x0, 0.05)), Inches(0.28))
            bar.fill.solid()
            bar.fill.fore_color.rgb = BLUE
            bar.line.fill.background()
            xb = chart_l + (r.base - vmin) / span * chart_w
            tick = sl.shapes.add_shape(1, Inches(xb), Inches(y - 0.05),
                                       Emu(int(914400 * 0.03)), Inches(0.38))
            tick.fill.solid()
            tick.fill.fore_color.rgb = INK
            tick.line.fill.background()
            y += 0.6
        xp = chart_l + (price - vmin) / span * chart_w
        pl = sl.shapes.add_shape(1, Inches(xp), Inches(1.2),
                                 Emu(int(914400 * 0.025)), Inches(y - 1.5))
        pl.fill.solid()
        pl.fill.fore_color.rgb = ROSE
        pl.line.fill.background()
        text(sl, xp - 0.7, y + 0.05, 2.2, 0.4,
             f"price {_money(price, cur, 0)}", 11, True, ROSE)

    # --- slide 4: comps ----------------------------------------------------------- #
    comps = report.comps
    if comps and comps.peers:
        sl = add_slide()
        text(sl, 0.7, 0.4, 12, 0.6, "Trading comps", 24, True)
        rows_c = [comps.target] + list(comps.peers)
        tbl = sl.shapes.add_table(
            len(rows_c) + 1, 6, Inches(0.7), Inches(1.2), Inches(11.9),
            Inches(0.35 * (len(rows_c) + 1))).table
        for j, h in enumerate(["Ticker", "Mkt cap", "EV/EBITDA", "EV/Sales", "P/E", "P/B"]):
            tbl.cell(0, j).text = h
        for i, rrow in enumerate(rows_c, start=1):
            tbl.cell(i, 0).text = rrow.ticker + (" (target)" if i == 1 else "")
            tbl.cell(i, 1).text = (
                f"{rrow.market_cap / 1e9:,.0f}B" if rrow.market_cap else "n/a")
            for j, attr in enumerate(["ev_ebitda", "ev_sales", "pe", "pb"], start=2):
                v = getattr(rrow, attr)
                tbl.cell(i, j).text = f"{v:.1f}x" if v else "n/a"

    # --- slide 5: thesis / risks (AI note) ------------------------------------------ #
    if note:
        sl = add_slide()
        text(sl, 0.7, 0.4, 12, 0.6, "Thesis & risks", 24, True)
        text(sl, 0.7, 1.1, 5.8, 0.5, "Thesis", 16, True, GREEN)
        bullets(sl, 0.7, 1.6, 5.8, 5.0, (note.get("thesis") or [])[:6])
        text(sl, 6.9, 1.1, 5.8, 0.5, "Risks", 16, True, ROSE)
        bullets(sl, 6.9, 1.6, 5.8, 5.0, (note.get("risks") or [])[:6])
        cats = note.get("catalysts") or []
        if cats:
            text(sl, 0.7, 5.9, 5.8, 0.4, "Catalysts", 14, True, AMBER)
            bullets(sl, 0.7, 6.3, 11.9, 1.0, cats[:3], size=11)

    # --- footer slide ------------------------------------------------------------ #
    sl = add_slide()
    text(sl, 0.7, 3.0, 11.9, 1.2,
         "Generated by Equity Research Automation — model estimates driven by user "
         "assumptions. Research support, not investment advice.", 14, False, DIM)

    path = _ensure_out() / f"{_safe(ticker.upper())}_briefing.pptx"
    prs.save(str(path))
    return str(path)
